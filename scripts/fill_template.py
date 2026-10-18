"""Fill the Vyapaar import template from a batch JSON file and move the processed images.

Usage: python scripts/fill_template.py <batch.json>

The batch file is a list of rows. Each row has an "image" (file name in the input
folder) and values keyed by the exact template headers, for example
{"image": "IMG_1.jpg", "Item name*": "...", "Sale price": 35}.
Rows without an "Item code" get a new internal 10-digit code starting with 9.
"""

import io
import json
import re
import shutil
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent.parent
TEMPLATE = REPO / "resources" / "Import Items Template.xlsx"
SCANNING = Path(r"C:\Users\DELL\Documents\scanning")
COLUMNS = 21  # A:U; column AE holds the category dropdown list and is kept.
# Codes starting with 2 are read as weighing-scale barcodes (weight/price), so avoid them.
CODE_PREFIX = "9"
CODE_LENGTH = 10


def load_template():
    # The template has an invalid "undefined" validation range that openpyxl rejects.
    buffer = io.BytesIO()
    with zipfile.ZipFile(TEMPLATE) as src, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item)
            if item.filename.startswith("xl/worksheets/"):
                data = re.sub(
                    rb'<dataValidation [^>]*sqref="undefined[^"]*".*?</dataValidation>',
                    b"", data, flags=re.S,
                )
            dst.writestr(item, data)
    return load_workbook(buffer)


def last_internal_number():
    highest = 0
    for path in (SCANNING / "output").glob("*.xlsx"):
        workbook = load_workbook(path, read_only=True)
        rows = workbook["Item Details"].iter_rows(min_row=2, min_col=2, max_col=2, values_only=True)
        for (code,) in rows:
            code = str(code or "")
            if len(code) == CODE_LENGTH and code.isdigit() and code.startswith(CODE_PREFIX):
                highest = max(highest, int(code[len(CODE_PREFIX):]))
        workbook.close()
    return highest


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    rows = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8-sig"))

    workbook = load_template()
    sheet = workbook["Item Details"]
    headers = [sheet.cell(1, column).value for column in range(1, COLUMNS + 1)]

    for number, row in enumerate(rows, 1):
        unknown = set(row) - set(headers) - {"image"}
        if unknown:
            sys.exit(f"Row {number}: unknown columns {sorted(unknown)}")
        if not row.get("Item name*"):
            sys.exit(f"Row {number}: Item name* is required")
        if not (SCANNING / "input" / row.get("image", "")).is_file():
            sys.exit(f"Row {number}: image not found in input: {row.get('image')}")
        if (SCANNING / "processed" / row["image"]).exists():
            sys.exit(f"Row {number}: image already exists in processed: {row['image']}")

    # Remove the template's sample rows and notes, including their borders and colours.
    for row_index in range(2, sheet.max_row + 1):
        for column in range(1, COLUMNS + 1):
            cell = sheet.cell(row_index, column)
            cell.value = None
            cell.style = "Normal"
    workbook["Batch Details"].delete_rows(2, workbook["Batch Details"].max_row)

    next_number = last_internal_number() + 1
    for row_index, row in enumerate(rows, 2):
        if not row.get("Item code"):
            row["Item code"] = f"{CODE_PREFIX}{next_number:0{CODE_LENGTH - len(CODE_PREFIX)}d}"
            next_number += 1
        for column, header in enumerate(headers, 1):
            value = row.get(header)
            if value is not None and value != "":
                cell = sheet.cell(row_index, column, value)
                if header == "Item code":
                    cell.value = str(value)
                    cell.number_format = "@"

    output = SCANNING / "output" / f"inventory-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    workbook.save(output)
    for row in rows:
        shutil.move(SCANNING / "input" / row["image"], SCANNING / "processed" / row["image"])

    print(f"Saved {len(rows)} rows to {output}")
    for row in rows:
        print(f"  {row['image']}: {row['Item code']}  {row['Item name*']}")


if __name__ == "__main__":
    main()
