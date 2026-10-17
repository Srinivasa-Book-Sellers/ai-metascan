"""Sample test case: 4 images in tests/sample/input, the confirmed rows in
tests/sample/batch.json, and the expected Excel in tests/sample/expected-output.xlsx.

Run from the repository root: python -m unittest discover -s tests
"""

import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

SAMPLE = Path(__file__).parent / "sample"
sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))
import fill_template  # noqa: E402


def item_rows(path):
    workbook = load_workbook(path, read_only=True)
    rows = [
        row for row in workbook["Item Details"].iter_rows(min_row=2, max_col=21, values_only=True)
        if any(value is not None for value in row)
    ]
    workbook.close()
    return rows


class SampleTest(unittest.TestCase):
    def test_sample_batch_matches_expected_output(self):
        scanning = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, scanning)
        for folder in ("output", "processed", "records"):
            (scanning / folder).mkdir()
        shutil.copytree(SAMPLE / "input", scanning / "input")

        fill_template.SCANNING = scanning
        sys.argv = ["fill_template.py", str(SAMPLE / "batch.json")]
        fill_template.main()

        output = next((scanning / "output").glob("inventory-*.xlsx"))
        self.assertEqual(item_rows(output), item_rows(SAMPLE / "expected-output.xlsx"))
        self.assertEqual(len(item_rows(output)), 4)
        self.assertEqual(
            sorted(path.name for path in (scanning / "processed").iterdir()),
            sorted(path.name for path in (SAMPLE / "input").iterdir()),
        )
        self.assertEqual(list((scanning / "input").iterdir()), [])


if __name__ == "__main__":
    unittest.main()
